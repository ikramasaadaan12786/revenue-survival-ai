import os
import asyncio
from dotenv import load_dotenv

load_dotenv()

async def generate_mission_1013_excel():
    from app.core.database import AsyncSessionLocal
    from app.services.intelligence.daily_excel_service import daily_excel_service
    
    async with AsyncSessionLocal() as session:
        buffer, filename, summary = await daily_excel_service.build_daily_workbook(
            session=session,
            mission_id=1013
        )
        reports_dir = os.path.join(os.path.dirname(__file__), "reports")
        os.makedirs(reports_dir, exist_ok=True)
        file_path = os.path.join(reports_dir, filename)
        with open(file_path, "wb") as f:
            f.write(buffer.getvalue())
        print(f"Successfully generated clean Excel workbook for Mission #1013 at {file_path}")
        print("Summary metrics:", summary)
        
        # Verify sheets
        import openpyxl
        wb = openpyxl.load_workbook(file_path)
        print("Sheet names in workbook:", wb.sheetnames)
        for name in wb.sheetnames:
            ws = wb[name]
            print(f" - Sheet '{name}': {ws.max_row} rows (including header), {ws.max_column} cols")

if __name__ == "__main__":
    asyncio.run(generate_mission_1013_excel())
